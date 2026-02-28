import os
import fitz  # PyMuPDF
from docx import Document as DocxDocument
import openpyxl

class PDFExtractor:
    @staticmethod
    def extract(file_path: str) -> str:
        text = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text.append(page.get_text())
        return "\n\n".join(text)

class DocxExtractor:
    @staticmethod
    def extract(file_path: str) -> str:
        doc = DocxDocument(file_path)
        return "\n".join([para.text for para in doc.paragraphs])

class XlsxExtractor:
    @staticmethod
    def extract(file_path: str) -> str:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            text.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) for cell in row if cell is not None])
                if row_text.strip():
                    text.append(row_text)
        return "\n".join(text)

class TxtExtractor:
    @staticmethod
    def extract(file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

def get_extractor(file_path: str):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return PDFExtractor()
    elif ext == '.docx':
        return DocxExtractor()
    elif ext == '.xlsx':
        return XlsxExtractor()
    elif ext == '.txt':
        return TxtExtractor()
    else:
        raise ValueError(f"Unsupported file format: {ext}")
